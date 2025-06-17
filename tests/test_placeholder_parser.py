"""
Unit tests for the PlaceholderParser class.
"""

import unittest
from unittest.mock import patch, MagicMock
from docx import Document
from docx.shared import Inches

from report_compiler.document.placeholder_parser import PlaceholderParser


class TestPlaceholderParser(unittest.TestCase):
    """Test cases for PlaceholderParser class."""
    
    def setUp(self):
        """Set up test fixtures."""
        self.parser = PlaceholderParser()
    
    def test_parse_overlay_placeholder_basic(self):
        """Test parsing basic overlay placeholder."""
        text = "[[OVERLAY: appendices/sketch.pdf]]"
        result = self.parser.parse_overlay_placeholder(text)
        
        expected = {
            'pdf_path': 'appendices/sketch.pdf',
            'page_spec': None,
            'crop': True
        }
        self.assertEqual(result, expected)
    
    def test_parse_overlay_placeholder_with_page(self):
        """Test parsing overlay placeholder with page specification."""
        text = "[[OVERLAY: diagrams/chart.pdf, page=2]]"
        result = self.parser.parse_overlay_placeholder(text)
        
        expected = {
            'pdf_path': 'diagrams/chart.pdf',
            'page_spec': '2',
            'crop': True
        }
        self.assertEqual(result, expected)
    
    def test_parse_overlay_placeholder_with_crop_false(self):
        """Test parsing overlay placeholder with crop=false."""
        text = "[[OVERLAY: full_page.pdf, crop=false]]"
        result = self.parser.parse_overlay_placeholder(text)
        
        expected = {
            'pdf_path': 'full_page.pdf',
            'page_spec': None,
            'crop': False
        }
        self.assertEqual(result, expected)
    
    def test_parse_overlay_placeholder_with_page_and_crop(self):
        """Test parsing overlay placeholder with both page and crop parameters."""
        text = "[[OVERLAY: report.pdf, page=1-3, crop=false]]"
        result = self.parser.parse_overlay_placeholder(text)
        
        expected = {
            'pdf_path': 'report.pdf',
            'page_spec': '1-3',
            'crop': False
        }
        self.assertEqual(result, expected)
    
    def test_parse_overlay_placeholder_with_spaces(self):
        """Test parsing overlay placeholder with extra spaces."""
        text = "[[OVERLAY:  appendices/sketch.pdf , page = 2 , crop = false ]]"
        result = self.parser.parse_overlay_placeholder(text)
        
        expected = {
            'pdf_path': 'appendices/sketch.pdf',
            'page_spec': '2',
            'crop': False
        }
        self.assertEqual(result, expected)
    
    def test_parse_overlay_placeholder_invalid(self):
        """Test parsing invalid overlay placeholder."""
        text = "[[INVALID: not_a_placeholder]]"
        result = self.parser.parse_overlay_placeholder(text)
        self.assertIsNone(result)
    
    def test_parse_insert_placeholder_basic(self):
        """Test parsing basic insert placeholder."""
        text = "[[INSERT: appendices/analysis.pdf]]"
        result = self.parser.parse_insert_placeholder(text)
        
        expected = {
            'pdf_path': 'appendices/analysis.pdf',
            'page_spec': None
        }
        self.assertEqual(result, expected)
    
    def test_parse_insert_placeholder_with_pages(self):
        """Test parsing insert placeholder with page specification."""
        text = "[[INSERT: report.pdf:1-5]]"
        result = self.parser.parse_insert_placeholder(text)
        
        expected = {
            'pdf_path': 'report.pdf',
            'page_spec': '1-5'
        }
        self.assertEqual(result, expected)
    
    def test_parse_insert_placeholder_complex_pages(self):
        """Test parsing insert placeholder with complex page specification."""
        text = "[[INSERT: calculations.pdf:2-4,7,9-]]"
        result = self.parser.parse_insert_placeholder(text)
        
        expected = {
            'pdf_path': 'calculations.pdf',
            'page_spec': '2-4,7,9-'
        }
        self.assertEqual(result, expected)
    
    def test_parse_insert_placeholder_with_spaces(self):
        """Test parsing insert placeholder with extra spaces."""
        text = "[[ INSERT : appendices/analysis.pdf : 1-3 ]]"
        result = self.parser.parse_insert_placeholder(text)
        
        expected = {
            'pdf_path': 'appendices/analysis.pdf',
            'page_spec': '1-3'
        }
        self.assertEqual(result, expected)
    
    def test_parse_insert_placeholder_invalid(self):
        """Test parsing invalid insert placeholder."""
        text = "[[INVALID: not_a_placeholder]]"
        result = self.parser.parse_insert_placeholder(text)
        self.assertIsNone(result)
    
    def test_resolve_relative_path_relative(self):
        """Test resolving relative path."""
        doc_path = "/documents/report.docx"
        pdf_path = "appendices/sketch.pdf"
        result = self.parser.resolve_relative_path(pdf_path, doc_path)
        expected = "/documents/appendices/sketch.pdf"
        self.assertEqual(result, expected)
    
    def test_resolve_relative_path_absolute(self):
        """Test resolving absolute path."""
        doc_path = "/documents/report.docx"
        pdf_path = "/shared/diagrams/chart.pdf"
        result = self.parser.resolve_relative_path(pdf_path, doc_path)
        expected = "/shared/diagrams/chart.pdf"
        self.assertEqual(result, expected)
    
    def test_resolve_relative_path_windows_absolute(self):
        """Test resolving Windows absolute path."""
        doc_path = "C:\\Documents\\report.docx"
        pdf_path = "C:\\Shared\\chart.pdf"
        result = self.parser.resolve_relative_path(pdf_path, doc_path)
        expected = "C:\\Shared\\chart.pdf"
        self.assertEqual(result, expected)
    
    @patch('docx.Document')
    def test_find_placeholders_in_document(self, mock_document):
        """Test finding placeholders in document."""
        # Create mock document structure
        mock_doc = MagicMock()
        mock_document.return_value = mock_doc
        
        # Mock paragraphs
        mock_para1 = MagicMock()
        mock_para1.text = "Normal paragraph text"
        mock_para2 = MagicMock()
        mock_para2.text = "[[INSERT: appendices/analysis.pdf:1-3]]"
        
        # Mock tables
        mock_table = MagicMock()
        mock_row = MagicMock()
        mock_cell = MagicMock()
        mock_cell.text = "[[OVERLAY: sketches/diagram.pdf, page=2]]"
        mock_row.cells = [mock_cell]
        mock_table.rows = [mock_row]
        
        mock_doc.paragraphs = [mock_para1, mock_para2]
        mock_doc.tables = [mock_table]
        
        # Mock path resolution
        with patch.object(self.parser, 'resolve_relative_path') as mock_resolve:
            mock_resolve.side_effect = lambda pdf, doc: f"/resolved/{pdf}"
            
            placeholders = self.parser.find_placeholders_in_document("/test/report.docx")
        
        # Verify results
        self.assertEqual(len(placeholders), 2)
        
        # Check INSERT placeholder
        insert_ph = next(ph for ph in placeholders if ph['type'] == 'INSERT')
        self.assertEqual(insert_ph['pdf_path'], '/resolved/appendices/analysis.pdf')
        self.assertEqual(insert_ph['page_spec'], '1-3')
        
        # Check OVERLAY placeholder
        overlay_ph = next(ph for ph in placeholders if ph['type'] == 'OVERLAY')
        self.assertEqual(overlay_ph['pdf_path'], '/resolved/sketches/diagram.pdf')
        self.assertEqual(overlay_ph['page_spec'], '2')
    
    def test_is_single_cell_table_true(self):
        """Test identifying single-cell table."""
        mock_table = MagicMock()
        mock_row = MagicMock()
        mock_row.cells = [MagicMock()]  # Single cell
        mock_table.rows = [mock_row]  # Single row
        
        result = self.parser.is_single_cell_table(mock_table)
        self.assertTrue(result)
    
    def test_is_single_cell_table_false_multiple_rows(self):
        """Test identifying table with multiple rows."""
        mock_table = MagicMock()
        mock_row1 = MagicMock()
        mock_row1.cells = [MagicMock()]
        mock_row2 = MagicMock()
        mock_row2.cells = [MagicMock()]
        mock_table.rows = [mock_row1, mock_row2]  # Multiple rows
        
        result = self.parser.is_single_cell_table(mock_table)
        self.assertFalse(result)
    
    def test_is_single_cell_table_false_multiple_cells(self):
        """Test identifying table with multiple cells."""
        mock_table = MagicMock()
        mock_row = MagicMock()
        mock_row.cells = [MagicMock(), MagicMock()]  # Multiple cells
        mock_table.rows = [mock_row]
        
        result = self.parser.is_single_cell_table(mock_table)
        self.assertFalse(result)


if __name__ == '__main__':
    unittest.main()
